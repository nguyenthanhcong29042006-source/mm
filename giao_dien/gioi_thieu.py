# -*- coding: utf-8 -*-
import streamlit as st
import os
import base64
import html
from pathlib import Path
from core import auth

ROOT = Path(__file__).resolve().parent.parent

# ==============================================================================
# TỐI ƯU HÓA KHOẢNG TRẮNG & RESPONSIVE CHO TRANG GIỚI THIỆU
# ==============================================================================
st.markdown("""
<style>
    /* Thu hẹp lề trên để nội dung nối tiếp mượt mà ngay bên dưới header chung của app.py */
    .block-container {
        padding-top: 0.5rem !important;
        padding-bottom: 3rem !important;
        max-width: 900px !important;
    }
    
    /* Tối ưu riêng cho màn hình điện thoại di động */
    @media screen and (max-width: 640px) {
        .block-container {
            padding-top: 0.3rem !important;
            padding-left: 0.8rem !important;
            padding-right: 0.8rem !important;
        }
    }

    /* Thiết kế lại nút st.page_link thành dạng minimalist sang trọng, gọn gàng */
    [data-testid="stPageLink"] {
        background-color: #fcfcfc !important;
        border: 1px solid #e2e8f0 !important;
        border-radius: 8px !important;
        padding: 6px 14px !important;
        transition: all 0.25s ease-in-out !important;
        width: fit-content !important;
        box-shadow: 0 1px 3px rgba(0,0,0,0.02) !important;
    }
    [data-testid="stPageLink"]:hover {
        background-color: #f1f5f9 !important;
        border-color: #cbd5e1 !important;
        box-shadow: 0 4px 12px rgba(0,0,0,0.06) !important;
        transform: translateY(-1px);
    }
    [data-testid="stPageLink"] span {
        font-family: 'Times New Roman', Times, serif !important;
        font-size: 15px !important;
        color: #003366 !important;
        font-weight: 600 !important;
    }

    /* Tối ưu thanh chọn ngôn ngữ góc trên bên trái */
    [data-testid="stSegmentedControl"] {
        display: inline-flex !important;
        justify-content: flex-start !important;
    }
</style>
""", unsafe_allow_html=True)

# Nút quay lại trang Hỏi đáp chính được tinh chỉnh gọn gàng, cao cấp
st.page_link("giao_dien/cong_dan.py", label="Quay lại trang Hỏi đáp chính", icon="⬅️")
st.markdown("<hr style='margin: 8px 0 15px 0;'>", unsafe_allow_html=True)

# ==============================================================================
# CƠ CHẾ LƯU TRỮ FILE CỨNG & QUẢN TRỊ NỘI DUNG GIỚI THIỆU
# ==============================================================================
DATA_DIR = ROOT / "data"
DATA_FILE = DATA_DIR / "gioi_thieu.md"

# Đường dẫn đến các file âm thanh thu sẵn trong thư mục audio
AUDIO_MONG_FILE = ROOT / "audio" / "gioi_thieu_mong.m4a"
AUDIO_VI_FILE = ROOT / "audio" / "gioi_thieu_vi.m4a"

def load_intro_content():
    """Đọc nội dung từ file cứng, nếu chưa có thì trả về nội dung chuẩn đầy đủ."""
    try:
        if DATA_FILE.exists():
            content = DATA_FILE.read_text(encoding="utf-8")
            if content.strip():
                return content
    except Exception as e:
        st.error(f"Lỗi đọc file: {e}")
        pass
    
    return """<h2 style="text-align: center; color: #003366;">GIỚI THIỆU DỰ ÁN</h2>
<h2 style="text-align: center; color: #003366;">LUẬT GẦN BẢN</h2>
<p style="text-align: center;"><b>Trợ lý thủ tục hành chính bằng giọng nói tiếng mẹ đẻ cho đồng bào dân tộc thiểu số</b></p>
<p style="text-align: center;"><i>“Chuyển đổi số: Không để ai bị bỏ lại phía sau”</i></p>
<hr>
<h3>I. Bối cảnh và bài toán xã hội</h3>
<p>Trong tiến trình chuyển đổi số quốc gia, hạ tầng công nghệ, điện lưới và điện thoại thông minh đã cơ bản phủ sóng đến các bản làng vùng cao. Tuy nhiên, một nghịch lý vẫn đang diễn ra tại bộ phận "Một cửa" của nhiều ủy ban nhân dân cấp xã: Người dân tộc thiểu số vẫn phải đi lại nhiều lần, thậm chí bỏ cuộc khi thực hiện các thủ tục hành chính thiết yếu như khai sinh, khai tử, kết hôn…</p>
<p>Đề án nhận diện nguyên nhân cốt lõi không nằm ở khoảng cách địa lý hay sự thiếu hụt thiết bị, mà nằm ở một rào cản vô hình mang tên "ngôn ngữ và chữ viết". Cụ thể, người dân đang phải đối diện với 3 lớp rào cản chồng lấn lên nhau:</p>
<ul>
    <li><b>Rào cản về ngôn ngữ:</b> Toàn bộ văn bản pháp luật và biểu mẫu đều viết bằng tiếng Việt, trong khi một bộ phận lớn đồng bào (đặc biệt là phụ nữ và người lớn tuổi) giao tiếp chủ yếu bằng tiếng mẹ đẻ (ví dụ: tiếng Mông).</li>
    <li><b>Rào cản về chữ viết:</b> Các giải pháp số hóa hiện tại đều mặc định người dùng có khả năng đọc hiểu trên màn hình, bỏ qua nhóm yếu thế đặc biệt là người dân tộc có tỷ lệ mù chữ cao.</li>
    <li><b>Rào cản về thuật ngữ:</b> Kể cả khi biết chữ, những thuật ngữ pháp lý chuyên ngành vẫn là một thách thức lớn đối với nhận thức của người dân.</li>
</ul>
<p>Vấn đề này không chỉ gây thiệt thòi cho người dân mà còn tạo áp lực khổng lồ lên đội ngũ cán bộ tư pháp - hộ tịch. Họ liên tục phải dành thời gian giải thích lặp đi lặp lại một quy định bằng lời nói, đối mặt với rủi ro hướng dẫn sai và lãng phí thời gian khi phải trả lại hồ sơ thiếu sót.</p>

<h3>II. Giải pháp Luật Gần Bản</h3>
<p>Đánh giá trên tình hình thực tiễn, thay vì yêu cầu người dân phải học chữ để hiểu luật, Luật Gần Bản đảo ngược cách tiếp cận: Buộc hệ thống công nghệ phải học cách nói tiếng của người dân.</p>
<p>Luật Gần Bản là một trợ lý ảo hỗ trợ tra cứu thủ tục hành chính, vận hành hoàn toàn bằng giọng nói và định vị đây là một "dự án âm thanh" chứ không phải dự án chữ viết.</p>
<p>Trên phiên bản sản phẩm tối thiểu (MVP) mà bạn đang tiếp cận, giao diện được thiết kế tối giản hóa tuyệt đối để ngay cả người không biết chữ cũng có thể sử dụng:</p>
<ul>
    <li><b>Bước 1. Lựa chọn:</b> Chọn ngôn ngữ được hiển thị trên màn hình (“Tiếng Mông” hoặc “Tiếng Việt”)</li>
    <li><b>Bước 2. Thao tác một chạm:</b> Người dùng nhấn vào biểu tượng Micro cỡ lớn ở trung tâm và nói ra nhu cầu của mình. (Ví dụ: “Tôi muốn làm giấy khai sinh cho con”).</li>
    <li><b>Bước 3: Tiếp nhận thông tin và trả lời câu hỏi:</b> Hệ thống tự động phân tích nhu cầu, tra cứu quy định pháp luật, đưa ra câu trả lời đã được đơn giản hóa bằng văn bản và bản audio tiếng dân tộc để hướng dẫn người dân.</li>
</ul>

<h3>III. Cơ chế vận hành và tính bảo đảm</h3>
<p>Đưa Trí tuệ nhân tạo (AI) vào lĩnh vực pháp luật luôn đi kèm rủi ro nghiêm trọng: Hệ thống có thể tự bịa đặt thông tin và hướng dẫn sai luật. Luật Gần Bản giải quyết rủi ro này bằng triết lý thiết kế: "Trí tuệ nhân tạo soạn thảo, con người ký duyệt". Dự án đảm bảo không có bất kỳ câu trả lời nào được phát ra cho người dân nếu chưa có một cán bộ đứng ra chịu trách nhiệm.</p>
<p>Dưới góc độ kỹ thuật, hệ thống được cấu trúc thành 6 khối độc lập. Trong đó, điểm khác biệt tạo nên sự an toàn và tính đảm bảo về mặt pháp lý nằm ở khối Quản trị và kiểm duyệt dành riêng cho đội ngũ cán bộ.</p>
<ul>
    <li>Khi có một thủ tục mới được đưa vào kho dữ liệu, máy tính sẽ tự động diễn giải quy định pháp luật thành ngôn ngữ đời thường (bản nháp bị “khóa”, tuyệt đối không phát ra cho người dân).</li>
    <li>Công chức tư pháp – hộ tịch sẽ truy cập vào Bảng điều khiển, đối chiếu bản gốc và bản máy soạn, trực tiếp chỉnh sửa câu chữ cho phù hợp với đặc thù địa phương rồi mới nhấn nút "Duyệt".</li>
    <li>Chỉ khi có lệnh duyệt này, âm thanh mới được phát tới người dùng, kèm theo lưu vết vĩnh viễn tên cán bộ kiểm duyệt và thời gian thực hiện.</li>
</ul>
<p>Cơ chế này đảm bảo mọi phát ngôn của hệ thống đều có một chủ thể là con người chịu trách nhiệm công vụ, không để máy móc tự ý phát ngôn nhân danh cơ quan nhà nước.</p>

<h3>IV. Tầm nhìn dài hạn</h3>
<p>Mục tiêu hiện tại của dự án là giúp người dân có sự chuẩn bị chính xác nhất trước khi bước ra khỏi nhà, giảm tỷ lệ hồ sơ phải làm lại. Trong tương lai, đề án hướng tới tích hợp sâu với hạ tầng số quốc gia nhằm tạo ra chu trình khép kín: người dân nói tiếng mẹ đẻ, hệ thống tự động điền biểu mẫu, đọc lại xác nhận và kết nối VNeID để nộp hồ sơ thẳng lên Cổng Dịch vụ công quốc gia.</p>
<p>Bằng sức mạnh của công nghệ vị nhân sinh, Luật Gần Bản cam kết hiện thực hóa thông điệp sâu sắc nhất của kỷ nguyên số: <i>“Chuyển đổi số: Không để ai bị bỏ lại phía sau”</i>.</p>
"""

def save_intro_content(content):
    """Lưu vĩnh viễn nội dung vào file cứng."""
    try:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        DATA_FILE.write_text(content, encoding="utf-8")
        return True
    except Exception as e:
        st.error(f"Lỗi khi lưu tệp hệ thống: {e}")
        return False

def docx_to_exact_html(docx_file) -> str:
    """Chuyển đổi file Docx thành mã HTML."""
    try:
        import docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        
        doc = docx.Document(docx_file)
        html_parts = []
        
        for p in doc.paragraphs:
            align_style = "text-align: left;"
            if p.alignment == WD_ALIGN_PARAGRAPH.CENTER:
                align_style = "text-align: center;"
            elif p.alignment == WD_ALIGN_PARAGRAPH.RIGHT:
                align_style = "text-align: right;"
            elif p.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:
                align_style = "text-align: justify;"
            
            p_content = []
            for run in p.runs:
                text = run.text
                if not text:
                    try:
                        drawings = run._r.xpath('.//a:blip')
                        for blip in drawings:
                            embed = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                            if embed and embed in doc.part.related_parts:
                                image_part = doc.part.related_parts[embed]
                                image_bytes = image_part.blob
                                b64_img = base64.b64encode(image_bytes).decode('utf-8')
                                p_content.append(f'<div style="text-align: center;"><img src="data:image/png;base64,{b64_img}" style="max-width: 100%; height: auto; margin: 20px auto; border-radius: 6px; box-shadow: 0 4px 12px rgba(0,0,0,0.15);" /></div>')
                    except Exception:
                        pass
                    continue
                
                safe_text = html.escape(text)
                style_runs = []
                if run.bold:
                    safe_text = f"<b>{safe_text}</b>"
                if run.italic:
                    safe_text = f"<i>{safe_text}</i>"
                if run.font.color and run.font.color.rgb:
                    rgb = run.font.color.rgb
                    hex_color = f"#{rgb[0]:02x}{rgb[1]:02x}{rgb[2]:02x}"
                    style_runs.append(f"color: {hex_color};")
                
                if style_runs:
                    style_str = " ".join(style_runs)
                    safe_text = f'<span style="{style_str}">{safe_text}</span>'
                
                p_content.append(safe_text)
            
            full_p_text = "".join(p_content)
            if full_p_text.strip() or '<img' in full_p_text:
                tag = "p"
                style_name = p.style.name.lower()
                if "heading 1" in style_name or "title" in style_name:
                    tag = "h2"
                elif "heading 2" in style_name:
                    tag = "h3"
                
                html_parts.append(f'<{tag} style="{align_style} margin-bottom: 10px;">{full_p_text}</{tag}>')
        
        for table in doc.tables:
            table_html = ['<div style="overflow-x: auto; margin: 15px 0;"><table style="border-collapse: collapse; width: 100%;">']
            for row in table.rows:
                table_html.append('<tr>')
                for cell in row.cells:
                    table_html.append(f'<td style="border: 1px solid #cccccc; padding: 8px 12px; text-align: left;">{html.escape(cell.text)}</td>')
                table_html.append('</tr>')
            table_html.append('</table></div>')
            html_parts.append("".join(table_html))
            
        return "\n".join(html_parts)
    except Exception as e:
        st.error(f"Không thể đọc file Word: {e}")
        return ""

if "intro_content" not in st.session_state:
    st.session_state["intro_content"] = load_intro_content()

# Phân quyền Admin chỉnh sửa nội dung
u = auth.nguoi_dang_nhap()
if u and u.get("vai_tro") == "admin":
    with st.expander("⚙️ BẢNG ĐIỀU KHIỂN ADMIN - CHỈNH SỬA NỘI DUNG", expanded=False):
        st.warning("Bạn đang đăng nhập bằng tài khoản Quản trị viên. Mọi thay đổi sẽ được lưu vĩnh viễn vào hệ thống.")
        
        tab_soan, tab_file = st.tabs(["✍️ Soạn thảo trực tiếp", "📁 Tải file lên (.md, .txt, .docx)"])
        
        with tab_soan:
            updated_content = st.text_area(
                "Soạn thảo nội dung giới thiệu (hỗ trợ HTML/Markdown):", 
                value=st.session_state["intro_content"], 
                height=300
            )
            if st.button("💾 Lưu nội dung soạn thảo", type="primary"):
                if save_intro_content(updated_content):
                    st.session_state["intro_content"] = updated_content
                    st.success("Đã lưu và cập nhật hệ thống thành công!")
                    st.rerun()
                
        with tab_file:
            st.caption("Tải lên file Markdown, Text hoặc Word (.docx). Hệ thống tự động bóc tách giữ nguyên định dạng gốc.")
            uploaded_file = st.file_uploader("Chọn file tài liệu tải lên", type=["md", "txt", "docx"])
            
            if uploaded_file is not None:
                file_content = ""
                if uploaded_file.name.endswith(".docx"):
                    file_content = docx_to_exact_html(uploaded_file)
                else:
                    try:
                        file_content = uploaded_file.read().decode("utf-8")
                    except Exception:
                        try:
                            file_content = uploaded_file.read().decode("latin-1")
                        except Exception as e:
                            st.error(f"Lỗi giải mã file: {e}")
                
                if file_content:
                    st.caption("Xem trước bố cục tài liệu:")
                    preview_html = f"""
                    <div style="background: #ffffff; color: #000000; padding: 20px; border-radius: 8px; box-shadow: 0 2px 10px rgba(0,0,0,0.1); font-family: 'Times New Roman', Times, serif; line-height: 1.7; max-height: 350px; overflow-y: auto; margin-bottom: 15px; box-sizing: border-box;">
                        {file_content}
                    </div>
                    """
                    st.markdown(preview_html, unsafe_allow_html=True)
                    
                    if st.button("🚀 Xác nhận cập nhật từ file", type="primary", use_container_width=True):
                        if save_intro_content(file_content):
                            st.session_state["intro_content"] = file_content
                            st.success("Đã cập nhật nội dung chuẩn định dạng từ file thành công!")
                            st.rerun()
                
    st.markdown("<hr style='margin: 10px 0;'>", unsafe_allow_html=True)

# ==============================================================================
# THANH CHỌN NGÔN NGỮ (GÓC TRÊN BÊN TRÁI) & PHÁT ÂM THANH THEO TÙY CHỌN
# ==============================================================================
col_lang, col_space = st.columns([3, 7])
with col_lang:
    selected_lang = st.segmented_control(
        "Chọn ngôn ngữ phát âm",
        options=["🔊 Tiếng Việt", "🔊 Tiếng Mông"],
        default=None,  # Không chọn sẵn, tránh tự động phát âm thanh khi mới vào trang
        key="intro_language_selector",
        label_visibility="collapsed"
    )

# Văn bản hiển thị chính luôn giữ nguyên tiếng Việt chuẩn
display_content = st.session_state["intro_content"]

# Xử lý phát âm thanh chính xác khi người dùng chủ động bấm chọn
if selected_lang == "🔊 Tiếng Mông":
    if AUDIO_MONG_FILE.exists():
        st.audio(str(AUDIO_MONG_FILE), format="audio/mp4", autoplay=True)
    else:
        st.warning("⚠️ Đang cập nhật tệp âm thanh tiếng Mông tại thư mục `audio/gioi_thieu_mong.m4a`.")
elif selected_lang == "🔊 Tiếng Việt":
    if AUDIO_VI_FILE.exists():
        st.audio(str(AUDIO_VI_FILE), format="audio/mp4", autoplay=True)
    else:
        st.info("💡 Để nghe bản đọc tiếng Việt, hãy đặt file âm thanh vào thư mục `audio/gioi_thieu_vi.m4a`.")

st.markdown("<div style='height: 4px;'></div>", unsafe_allow_html=True)

# ==============================================================================
# HIỂN THỊ NỘI DUNG CHÍNH (Responsive mượt mà trên cả PC và Mobile)
# ==============================================================================
document_html = f"""
<div style="
    background: #ffffff;
    color: #111111;
    padding: clamp(15px, 3.5vw, 45px);
    margin: 5px auto;
    max-width: 900px;
    width: 100%;
    border-radius: 6px;
    box-shadow: 0 4px 20px rgba(0,0,0,0.08);
    font-family: 'Times New Roman', Times, serif;
    line-height: 1.7;
    font-size: clamp(15px, 1.8vw, 17px);
    box-sizing: border-box;
">
    <style>
        img {{
            max-width: 100% !important;
            height: auto !important;
            display: block !important;
            margin: 20px auto !important;
            border-radius: 6px !important;
            box-shadow: 0 4px 12px rgba(0,0,0,0.12) !important;
        }}
        h2, h3, h4 {{
            color: #003366 !important;
            font-family: 'Times New Roman', Times, serif !important;
            font-weight: bold !important;
            margin-top: 20px !important;
            margin-bottom: 10px !important;
            word-wrap: break-word;
        }}
        p, li {{
            margin-bottom: 12px !important;
            word-wrap: break-word;
        }}
        table {{
            border-collapse: collapse;
            width: 100%;
            margin: 15px 0;
            display: block;
            overflow-x: auto;
            white-space: nowrap;
        }}
        th, td {{
            border: 1px solid #cccccc;
            padding: 8px 12px;
            text-align: left;
        }}
    </style>
    {display_content}
</div>
"""

st.markdown(document_html, unsafe_allow_html=True)
